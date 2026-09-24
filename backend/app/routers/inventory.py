"""盘点模块。

流程：发起盘点（按筛选条件把设备圈进任务）→ 逐台标记「已盘 / 异常」→ 看差异 → 导出 → 结束盘点。

关于「扫码盘点」的实现取舍（重要）：
    浏览器要调摄像头必须先满足 secure context，也就是 HTTPS 或 localhost。
    而本系统是局域网 http://192.168.x.x:8080 —— 手机上 getUserMedia 直接被浏览器拦死，
    引任何前端扫码库都救不回来。
    所以这里的做法是：用手机自带相机 / 微信扫设备上已有那张二维码 → 落到设备详情页 →
    详情页顶部自动弹出「盘点条」（/context/{asset_id} 查出来的），一点就标记完。
    不需要摄像头权限、不需要 HTTPS、微信内置浏览器里也能用。
"""

from __future__ import annotations

import io
import json
import urllib.parse
from datetime import date, datetime
from typing import Any, Optional

from fastapi import APIRouter, Depends, HTTPException, Query, status
from fastapi.responses import Response
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from ..database import get_db
from ..models import (
    Asset,
    AssetEventType,
    AssetStatus,
    AuditAction,
    AuditTarget,
    DeviceType,
    InventoryItem,
    InventoryResult,
    InventoryTask,
    InventoryTaskStatus,
)
from ..schemas import (
    InventoryContext,
    InventoryItemOut,
    InventoryMarkRequest,
    InventoryMarkResult,
    InventoryScope,
    InventoryTaskCreate,
    InventoryTaskDetail,
    InventoryTaskOut,
)
from ..services.audit import asset_label, record_audit
from ..services.events import record_event
from ..services.filters import build_conditions
from ..services.serialize import to_asset_out

router = APIRouter(prefix="/api/inventory", tags=["inventory"])

EXPORT_COLUMNS = [
    ("asset_code", "IT资产编号"),
    ("finance_code", "财务编码"),
    ("device_type_name", "设备类型"),
    ("brand", "品牌"),
    ("model", "型号"),
    ("serial_number", "序列号SN"),
    ("status_label", "系统状态"),
    ("user_name", "账面使用人"),
    ("location", "账面位置"),
    ("result_label", "盘点结果"),
    ("operator", "盘点人"),
    ("checked_at", "盘点时间"),
    ("note", "盘点备注"),
]


# --------------------------------------------------------------------------- #
# 工具
# --------------------------------------------------------------------------- #
def _get_task_or_404(db: Session, task_id: int) -> InventoryTask:
    task = db.get(InventoryTask, task_id)
    if task is None:
        raise HTTPException(status_code=404, detail=f"盘点任务 #{task_id} 不存在")
    return task


def _load_scope(task: InventoryTask) -> dict[str, Any]:
    if not task.scope:
        return {}
    try:
        parsed = json.loads(task.scope)
        return parsed if isinstance(parsed, dict) else {}
    except (ValueError, TypeError):
        return {}


def scope_text(db: Session, scope: dict[str, Any]) -> str:
    """把筛选条件翻译成一句人话，用在任务列表上。"""
    if not scope:
        return "全部设备"

    asset_ids = scope.get("asset_ids")
    if asset_ids:
        return f"指定 {len(asset_ids)} 台设备"

    parts: list[str] = []

    type_ids = [int(x) for x in str(scope.get("device_type_id") or "").split(",") if x.strip().isdigit()]
    if type_ids:
        names = db.execute(
            select(DeviceType.name).where(DeviceType.id.in_(type_ids))
        ).scalars().all()
        parts.append("类型：" + "、".join(names) if names else f"类型 ID {type_ids}")

    statuses = [s.strip() for s in str(scope.get("status") or "").split(",") if s.strip()]
    if statuses:
        parts.append("状态：" + "、".join(AssetStatus.LABELS.get(s, s) for s in statuses))

    if scope.get("user_name"):
        parts.append(f"使用人：{scope['user_name']}")
    if scope.get("location"):
        parts.append(f"位置：{scope['location']}")
    if scope.get("keyword"):
        parts.append(f"关键字：{scope['keyword']}")

    return " · ".join(parts) if parts else "全部设备"


def _counts(db: Session, task_ids: list[int]) -> dict[int, dict[str, int]]:
    """一次查出多个任务的进度，避免任务列表 N+1。"""
    if not task_ids:
        return {}
    rows = db.execute(
        select(InventoryItem.task_id, InventoryItem.result, func.count(InventoryItem.id))
        .where(InventoryItem.task_id.in_(task_ids))
        .group_by(InventoryItem.task_id, InventoryItem.result)
    ).all()

    out: dict[int, dict[str, int]] = {tid: {"total": 0, "checked": 0, "pending": 0, "abnormal": 0} for tid in task_ids}
    for tid, result, cnt in rows:
        bucket = out.setdefault(int(tid), {"total": 0, "checked": 0, "pending": 0, "abnormal": 0})
        bucket["total"] += int(cnt)
        if result in bucket:
            bucket[result] += int(cnt)
        else:
            bucket[InventoryResult.PENDING] += int(cnt)
    return out


def serialize_task(task: InventoryTask, counts: dict[str, int], scope_desc: str) -> InventoryTaskOut:
    total = counts.get("total", 0)
    checked = counts.get("checked", 0)
    return InventoryTaskOut(
        id=task.id,
        name=task.name,
        status=task.status,
        status_label=InventoryTaskStatus.LABELS.get(task.status, task.status),
        note=task.note,
        scope=_load_scope(task),
        scope_text=scope_desc,
        created_by=task.created_by,
        created_at=task.created_at,
        closed_by=task.closed_by,
        closed_at=task.closed_at,
        total=total,
        checked=checked,
        pending=counts.get("pending", 0),
        abnormal=counts.get("abnormal", 0),
        progress=round(checked / total * 100, 1) if total else 0.0,
    )


def serialize_item(item: InventoryItem) -> InventoryItemOut:
    asset = item.asset
    snapshot = (item.snapshot_user, item.snapshot_location, item.snapshot_status)
    current = (asset.user_name, asset.location, asset.status) if asset else (None, None, None)
    changed = item.result != InventoryResult.PENDING and snapshot != current

    return InventoryItemOut(
        id=item.id,
        task_id=item.task_id,
        asset_id=item.asset_id,
        asset=to_asset_out(asset) if asset else None,  # type: ignore[arg-type]
        result=item.result,
        result_label=InventoryResult.LABELS.get(item.result, item.result),
        snapshot_user=item.snapshot_user,
        snapshot_location=item.snapshot_location,
        snapshot_status=item.snapshot_status,
        operator=item.operator,
        note=item.note,
        checked_at=item.checked_at,
        changed_since=changed,
    )


def _select_assets_for_scope(db: Session, scope: InventoryScope) -> list[Asset]:
    if scope.asset_ids:
        ids = [int(i) for i in scope.asset_ids]
        rows = db.execute(select(Asset).where(Asset.id.in_(ids))).scalars().all()
        # 保持调用方给的顺序，前端圈选完再发起盘点时看到的顺序不会乱跳
        order = {aid: idx for idx, aid in enumerate(ids)}
        return sorted(rows, key=lambda a: order.get(a.id, 0))

    conditions = build_conditions(
        keyword=scope.keyword,
        device_type_id=scope.device_type_id,
        status=scope.status,
        user_name=scope.user_name,
        location=scope.location,
    )
    stmt = select(Asset)
    if conditions:
        stmt = stmt.where(*conditions)
    return list(db.execute(stmt.order_by(Asset.asset_code.asc())).scalars().all())


# --------------------------------------------------------------------------- #
# 任务
# --------------------------------------------------------------------------- #
@router.get("/tasks", response_model=list[InventoryTaskOut], summary="盘点任务列表")
def list_tasks(
    db: Session = Depends(get_db),
    task_status: Optional[str] = Query(None, alias="status", description="open / closed"),
    limit: int = Query(100, ge=1, le=500),
):
    stmt = select(InventoryTask)
    if task_status:
        stmt = stmt.where(InventoryTask.status == task_status)
    rows = list(
        db.execute(stmt.order_by(InventoryTask.status.asc(), InventoryTask.id.desc()).limit(limit))
        .scalars()
        .all()
    )
    counts = _counts(db, [t.id for t in rows])
    return [serialize_task(t, counts.get(t.id, {}), scope_text(db, _load_scope(t))) for t in rows]


@router.post(
    "/tasks",
    response_model=InventoryTaskOut,
    status_code=status.HTTP_201_CREATED,
    summary="发起盘点（按筛选条件圈定设备）",
)
def create_task(payload: InventoryTaskCreate, db: Session = Depends(get_db)):
    assets = _select_assets_for_scope(db, payload.scope)
    if not assets:
        raise HTTPException(
            status_code=400,
            detail="这个筛选条件下一台设备都没有，盘点任务发起不了。换个范围再试。",
        )

    scope_dict = payload.scope.model_dump(exclude_none=True)
    task = InventoryTask(
        name=payload.name,
        status=InventoryTaskStatus.OPEN,
        scope=json.dumps(scope_dict, ensure_ascii=False),
        note=payload.note,
        created_by=payload.created_by,
    )
    db.add(task)
    db.flush()

    for asset in assets:
        db.add(
            InventoryItem(
                task_id=task.id,
                asset_id=asset.id,
                result=InventoryResult.PENDING,
                # 快照账面值：盘点结束后回过头看「当时账上写的是谁用的」，差异分析才有依据
                snapshot_user=asset.user_name,
                snapshot_location=asset.location,
                snapshot_status=asset.status,
            )
        )

    scope_desc = scope_text(db, scope_dict)
    record_audit(
        db,
        AuditAction.INVENTORY_CREATE,
        target_type=AuditTarget.INVENTORY,
        target_id=task.id,
        target_label=f"盘点任务：{task.name}",
        summary=f"发起盘点《{task.name}》，范围：{scope_desc}，共 {len(assets)} 台设备",
        operator=payload.created_by,
        detail={"task_id": task.id, "task_name": task.name, "scope": scope_dict,
                "scope_text": scope_desc, "total": len(assets)},
    )

    db.commit()
    db.refresh(task)
    counts = {"total": len(assets), "checked": 0, "pending": len(assets), "abnormal": 0}
    return serialize_task(task, counts, scope_desc)


@router.get("/tasks/{task_id}", response_model=InventoryTaskDetail, summary="盘点详情（含明细与差异）")
def get_task(
    task_id: int,
    db: Session = Depends(get_db),
    result: Optional[str] = Query(None, description="只看某种结果：pending / checked / abnormal"),
):
    task = _get_task_or_404(db, task_id)
    rows = list(
        db.execute(
            select(InventoryItem)
            .where(InventoryItem.task_id == task_id)
            .order_by(InventoryItem.id.asc())
        ).scalars().all()
    )

    items = [serialize_item(i) for i in rows]
    counts = {
        "total": len(items),
        "checked": sum(1 for i in items if i.result == InventoryResult.CHECKED),
        "pending": sum(1 for i in items if i.result == InventoryResult.PENDING),
        "abnormal": sum(1 for i in items if i.result == InventoryResult.ABNORMAL),
    }

    base = serialize_task(task, counts, scope_text(db, _load_scope(task)))
    diff = [i for i in items if i.result != InventoryResult.CHECKED]
    shown = [i for i in items if i.result == result] if result else items

    return InventoryTaskDetail(**base.model_dump(), items=shown, diff=diff)


@router.post("/tasks/{task_id}/items", response_model=InventoryItemOut, summary="补录设备进盘点范围")
def add_item(task_id: int, asset_id: int = Query(...), db: Session = Depends(get_db)):
    task = _get_task_or_404(db, task_id)
    if task.status != InventoryTaskStatus.OPEN:
        raise HTTPException(status_code=409, detail="任务已经结束了，不能再往里加设备")
    asset = db.get(Asset, asset_id)
    if asset is None:
        raise HTTPException(status_code=404, detail=f"资产 #{asset_id} 不存在")

    exists = db.execute(
        select(InventoryItem).where(
            InventoryItem.task_id == task_id, InventoryItem.asset_id == asset_id
        )
    ).scalar_one_or_none()
    if exists is not None:
        return serialize_item(exists)

    item = InventoryItem(
        task_id=task_id,
        asset_id=asset_id,
        result=InventoryResult.PENDING,
        snapshot_user=asset.user_name,
        snapshot_location=asset.location,
        snapshot_status=asset.status,
    )
    db.add(item)
    # 补录也会改变盘点的范围，所以同样留痕 —— 不然"这轮到底盘了哪些"事后没法还原
    record_audit(
        db,
        AuditAction.INVENTORY_MARK,
        target_type=AuditTarget.INVENTORY,
        target_id=task.id,
        target_label=f"盘点任务：{task.name}",
        summary=f"把 {asset.asset_code} 补录进盘点范围",
        detail={"task_id": task.id, "asset_id": asset.id, "asset_code": asset.asset_code,
                "kind": "add_item"},
    )
    db.commit()
    db.refresh(item)
    return serialize_item(item)


@router.post("/tasks/{task_id}/mark", response_model=InventoryMarkResult, summary="标记盘点结果")
def mark_item(task_id: int, payload: InventoryMarkRequest, db: Session = Depends(get_db)):
    task = _get_task_or_404(db, task_id)
    if task.status != InventoryTaskStatus.OPEN:
        raise HTTPException(status_code=409, detail="任务已经结束了，不能继续标记")

    item = db.execute(
        select(InventoryItem).where(
            InventoryItem.task_id == task_id, InventoryItem.asset_id == payload.asset_id
        )
    ).scalar_one_or_none()
    if item is None:
        raise HTTPException(
            status_code=404,
            detail=f"资产 #{payload.asset_id} 不在这轮盘点的范围里。"
            f"如果确实要盘它，先点「补录」把它加进来。",
        )

    item.result = payload.result
    item.operator = payload.operator
    item.note = payload.note
    # 「未盘」= 把标记撤销回到初始状态，这时候不该留下盘点时间
    item.checked_at = None if payload.result == InventoryResult.PENDING else datetime.now()

    asset = db.get(Asset, item.asset_id)
    if asset is not None:
        label = InventoryResult.LABELS.get(payload.result, payload.result)
        record_event(
            db,
            asset,
            AssetEventType.INVENTORY,
            operator=payload.operator,
            note=payload.note or f"盘点《{task.name}》：{label}",
            detail={
                "task_id": task.id,
                "task_name": task.name,
                "result": payload.result,
                "snapshot_location": item.snapshot_location,
                "snapshot_user": item.snapshot_user,
            },
        )

    record_audit(
        db,
        AuditAction.INVENTORY_MARK,
        target_type=AuditTarget.ASSET,
        target_id=item.asset_id,
        target_label=asset_label(asset) if asset is not None else f"#{item.asset_id}",
        summary=f"盘点《{task.name}》标记为「{InventoryResult.LABELS.get(payload.result, payload.result)}」"
        + (f"；{payload.note}" if payload.note else ""),
        operator=payload.operator,
        detail={"task_id": task.id, "task_name": task.name, "asset_id": item.asset_id,
                "result": payload.result, "kind": "mark"},
    )

    db.commit()
    db.refresh(item)

    rows = db.execute(select(InventoryItem).where(InventoryItem.task_id == task_id)).scalars().all()
    counts = {
        "total": len(rows),
        "checked": sum(1 for r in rows if r.result == InventoryResult.CHECKED),
        "pending": sum(1 for r in rows if r.result == InventoryResult.PENDING),
        "abnormal": sum(1 for r in rows if r.result == InventoryResult.ABNORMAL),
    }
    return InventoryMarkResult(
        item=serialize_item(item),
        task=serialize_task(task, counts, scope_text(db, _load_scope(task))),
    )


@router.post("/tasks/{task_id}/close", response_model=InventoryTaskOut, summary="结束盘点")
def close_task(
    task_id: int,
    db: Session = Depends(get_db),
    operator: Optional[str] = Query(None),
):
    task = _get_task_or_404(db, task_id)
    if task.status == InventoryTaskStatus.CLOSED:
        raise HTTPException(status_code=409, detail="这个任务已经结束过了")

    task.status = InventoryTaskStatus.CLOSED
    task.closed_by = (operator or "").strip() or None
    task.closed_at = datetime.now()

    counts = _counts(db, [task.id]).get(task.id, {})
    record_audit(
        db,
        AuditAction.INVENTORY_CLOSE,
        target_type=AuditTarget.INVENTORY,
        target_id=task.id,
        target_label=f"盘点任务：{task.name}",
        summary=f"结束盘点《{task.name}》：应盘 {counts.get('total', 0)} 台，"
        f"已盘 {counts.get('checked', 0)}，未盘 {counts.get('pending', 0)}，异常 {counts.get('abnormal', 0)}",
        operator=operator,
        detail={"task_id": task.id, "task_name": task.name, **counts},
    )

    db.commit()
    db.refresh(task)
    return serialize_task(task, _counts(db, [task.id]).get(task.id, {}), scope_text(db, _load_scope(task)))


@router.post("/tasks/{task_id}/reopen", response_model=InventoryTaskOut, summary="重新开启盘点")
def reopen_task(
    task_id: int,
    db: Session = Depends(get_db),
    operator: Optional[str] = Query(None),
):
    task = _get_task_or_404(db, task_id)
    if task.status == InventoryTaskStatus.OPEN:
        raise HTTPException(status_code=409, detail="这个任务本来就是进行中")

    task.status = InventoryTaskStatus.OPEN
    task.closed_by = None
    task.closed_at = None
    record_audit(
        db,
        AuditAction.INVENTORY_REOPEN,
        target_type=AuditTarget.INVENTORY,
        target_id=task.id,
        target_label=f"盘点任务：{task.name}",
        summary=f"重新开启盘点《{task.name}》",
        operator=operator,
        detail={"task_id": task.id, "task_name": task.name},
    )
    db.commit()
    db.refresh(task)
    return serialize_task(task, _counts(db, [task.id]).get(task.id, {}), scope_text(db, _load_scope(task)))


@router.delete("/tasks/{task_id}", summary="删除盘点任务")
def delete_task(
    task_id: int,
    db: Session = Depends(get_db),
    operator: Optional[str] = Query(None),
):
    task = _get_task_or_404(db, task_id)
    name = task.name
    counts = _counts(db, [task.id]).get(task.id, {})
    record_audit(
        db,
        AuditAction.INVENTORY_DELETE,
        target_type=AuditTarget.INVENTORY,
        target_id=task_id,
        target_label=f"盘点任务：{name}",
        summary=f"删除盘点任务《{name}》（共 {counts.get('total', 0)} 条明细）",
        operator=operator,
        detail={"task_id": task_id, "task_name": name, **counts},
    )
    db.delete(task)
    db.commit()
    return {"ok": True, "deleted_id": task_id, "name": name}


# --------------------------------------------------------------------------- #
# 扫码上下文：手机上打开设备详情页时用
# --------------------------------------------------------------------------- #
@router.get(
    "/context/{asset_id}",
    response_model=Optional[InventoryContext],
    summary="该设备当前被哪个进行中的盘点任务盯着",
)
def get_context(asset_id: int, db: Session = Depends(get_db)):
    row = db.execute(
        select(InventoryItem, InventoryTask)
        .join(InventoryTask, InventoryTask.id == InventoryItem.task_id)
        .where(
            InventoryItem.asset_id == asset_id,
            InventoryTask.status == InventoryTaskStatus.OPEN,
        )
        .order_by(InventoryTask.id.desc())
    ).first()

    if row is None:
        return None

    item, task = row
    return InventoryContext(
        task_id=task.id,
        task_name=task.name,
        task_status=task.status,
        item_id=item.id,
        result=item.result,
        result_label=InventoryResult.LABELS.get(item.result, item.result),
        operator=item.operator,
        note=item.note,
        checked_at=item.checked_at,
    )


# --------------------------------------------------------------------------- #
# 导出
# --------------------------------------------------------------------------- #
def _export_rows(items: list[InventoryItemOut]) -> list[list[str]]:
    rows: list[list[str]] = []
    for item in items:
        asset = item.asset
        record = {
            "asset_code": asset.asset_code if asset else "",
            "finance_code": (asset.finance_code if asset else "") or "",
            "device_type_name": (asset.device_type_name if asset else "") or "",
            "brand": (asset.brand if asset else "") or "",
            "model": (asset.model if asset else "") or "",
            "serial_number": (asset.serial_number if asset else "") or "",
            "status_label": (asset.status_label if asset else "") or "",
            "user_name": (asset.user_name if asset else "") or "",
            "location": (asset.location if asset else "") or "",
            "result_label": item.result_label,
            "operator": item.operator or "",
            "checked_at": item.checked_at.strftime("%Y-%m-%d %H:%M") if item.checked_at else "",
            "note": item.note or "",
        }
        rows.append([str(record.get(key, "")) for key, _ in EXPORT_COLUMNS])
    return rows


def _content_disposition(filename: str) -> str:
    """中文文件名必须走 RFC 5987 的 filename*，否则浏览器下载下来是一串乱码。"""
    quoted = urllib.parse.quote(filename)
    return f"attachment; filename=\"inventory\"; filename*=UTF-8''{quoted}"


@router.get("/tasks/{task_id}/export", summary="导出盘点结果（CSV / xlsx）")
def export_task(
    task_id: int,
    db: Session = Depends(get_db),
    fmt: str = Query("xlsx", pattern="^(xlsx|csv)$", description="xlsx 或 csv"),
    only_diff: bool = Query(False, description="只导出差异（未盘 + 异常）"),
    result: Optional[str] = Query(None, description="只导出某种结果"),
):
    task = _get_task_or_404(db, task_id)
    rows = list(
        db.execute(
            select(InventoryItem)
            .where(InventoryItem.task_id == task_id)
            .order_by(InventoryItem.id.asc())
        ).scalars().all()
    )
    items = [serialize_item(i) for i in rows]
    if only_diff:
        items = [i for i in items if i.result != InventoryResult.CHECKED]
    if result:
        items = [i for i in items if i.result == result]

    safe_name = "".join(c for c in task.name if c not in '\\/:*?"<>|').strip() or f"盘点{task_id}"
    stamp = date.today().isoformat()
    suffix = "_差异" if only_diff else ""
    headers = [label for _, label in EXPORT_COLUMNS]
    data = _export_rows(items)

    if fmt == "csv":
        buf = io.StringIO()
        def esc(v: str) -> str:
            return '"' + v.replace('"', '""') + '"'
        buf.write(",".join(esc(h) for h in headers) + "\r\n")
        for row in data:
            buf.write(",".join(esc(v) for v in row) + "\r\n")
        # BOM 不能省：Excel 读 UTF-8 CSV 时没 BOM 就是一片乱码
        payload = "\ufeff" + buf.getvalue()
        return Response(
            content=payload.encode("utf-8"),
            media_type="text/csv; charset=utf-8",
            headers={"Content-Disposition": _content_disposition(f"盘点_{safe_name}{suffix}_{stamp}.csv")},
        )

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = Workbook()
    ws = wb.active
    ws.title = "盘点结果"

    head_fill = PatternFill("solid", fgColor="EEF3FF")
    head_font = Font(bold=True, color="1F2329")
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = head_fill
        cell.font = head_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in data:
        ws.append(row)

    widths = [16, 16, 10, 10, 22, 18, 10, 12, 22, 10, 10, 18, 24]
    for idx, width in enumerate(widths[: len(headers)], start=1):
        ws.column_dimensions[ws.cell(row=1, column=idx).column_letter].width = width
    ws.freeze_panes = "A2"

    # 汇总放第二个 sheet，交付给别人的时候一眼能看懂结论
    summary = wb.create_sheet("汇总")
    counts = _counts(db, [task_id]).get(task_id, {})
    summary_rows = [
        ("盘点任务", task.name),
        ("盘点范围", scope_text(db, _load_scope(task))),
        ("任务状态", InventoryTaskStatus.LABELS.get(task.status, task.status)),
        ("发起人", task.created_by or ""),
        ("发起时间", task.created_at.strftime("%Y-%m-%d %H:%M") if task.created_at else ""),
        ("结束时间", task.closed_at.strftime("%Y-%m-%d %H:%M") if task.closed_at else "未结束"),
        ("应盘总数", counts.get("total", 0)),
        ("已盘", counts.get("checked", 0)),
        ("未盘", counts.get("pending", 0)),
        ("异常", counts.get("abnormal", 0)),
        ("导出时间", datetime.now().strftime("%Y-%m-%d %H:%M")),
    ]
    for key, value in summary_rows:
        summary.append([key, value])
    summary.column_dimensions["A"].width = 14
    summary.column_dimensions["B"].width = 44
    for cell in summary["A"]:
        cell.font = Font(bold=True)

    out = io.BytesIO()
    wb.save(out)
    return Response(
        content=out.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": _content_disposition(f"盘点_{safe_name}{suffix}_{stamp}.xlsx")},
    )
